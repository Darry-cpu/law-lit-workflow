"""核验 .docx 里的脚注是否为 Word 原生页下注，以及文档能否正常打开。

用法::

    python docxcheck.py 文档.docx

为什么需要它：脚注最容易"看起来有、实际没有"——部件缺失时 Word 仍能打开，
只是把脚注引用吞掉或报"内容有问题"。所以要逐一确认四件事：部件本体、
内容类型声明、关系项、正文引用，四者缺一不可。
"""

import argparse
import pathlib
import re
import sys
import zipfile

import docx


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("docx")
    a = ap.parse_args()
    path = pathlib.Path(a.docx)

    print(f"文件：{path.name}  ({path.stat().st_size / 1024:.0f} KB)")

    ok = True
    with zipfile.ZipFile(path) as z:
        names = z.namelist()

        if "word/footnotes.xml" not in names:
            print("✗ 缺少 word/footnotes.xml —— 脚注不是真脚注")
            sys.exit(1)
        ft = z.read("word/footnotes.xml").decode("utf-8")
        notes = re.findall(r'<w:footnote w:id="([1-9]\d*)"', ft)
        seps = re.findall(r'<w:footnote w:type="(\w+)"', ft)
        doc = z.read("word/document.xml").decode("utf-8")
        refs = re.findall(r'<w:footnoteReference w:id="(\d+)"', doc)
        ct = z.read("[Content_Types].xml").decode("utf-8")
        rels = z.read("word/_rels/document.xml.rels").decode("utf-8")

        print(f"分离符脚注（分隔线）：{seps}")
        print(f"脚注条目数：{len(notes)}")
        print(f"正文中的脚注引用数：{len(refs)}")
        print(f"Content_Types 声明 footnotes+xml：{'footnotes+xml' in ct}")
        print(f"document.xml.rels 含 footnotes 关系：{'footnotes.xml' in rels}")

        if not ("footnotes+xml" in ct):
            print("✗ 内容类型未声明")
            ok = False
        if "footnotes.xml" not in rels:
            print("✗ 关系项缺失")
            ok = False
        if len(notes) != len(refs):
            print(f"✗ 脚注条目与正文引用数量不一致（{len(notes)} vs {len(refs)}）")
            ok = False

    d = docx.Document(str(path))
    paras = [p for p in d.paragraphs if p.text.strip()]
    print(f"\n可正常打开：是；非空段落数：{len(paras)}")
    print(f"首段：{paras[0].text[:50]}")
    print(f"末段：{paras[-1].text[:60]}")

    heads = [p.text for p in d.paragraphs if p.style.name.startswith("Heading")]
    print(f"\n标题层级（{len(heads)} 个）：")
    for h in heads:
        print("  " + h[:60])

    if ok:
        print("\n✓ 脚注为 Word 原生页下注，四要件齐备。")
    else:
        sys.exit(1)


if __name__ == "__main__":
    main()
