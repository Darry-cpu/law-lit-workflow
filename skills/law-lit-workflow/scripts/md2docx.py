"""把本项目的论文 Markdown 转成带**真脚注**（Word 页下注）的 .docx。

用法::

    python md2docx.py 输入.md 输出.docx

为什么不能直接用 python-docx 了事：python-docx 至今不提供脚注 API。Word 的脚注
不是普通文字，而是独立的 word/footnotes.xml 部件，需要同时改四样东西——
部件本体、[Content_Types].xml 的 Override、document.xml.rels 的关系项，
以及正文里的 w:footnoteReference 引用。缺任何一样，Word 打开时会报"内容有问题"
或把脚注整段吞掉。本脚本把这四步一次做齐。

排版按中文法学论文惯例：A4、正文宋体小四、1.5 倍行距、首行缩进 2 字符，
标题黑体，脚注小五号。
"""

import argparse
import pathlib
import re
import shutil
import zipfile

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

BODY_FONT = "宋体"
HEAD_FONT = "黑体"
BODY_SIZE = Pt(12)       # 小四
FOOT_SIZE = Pt(9)        # 小五
LINE_SPACING = 1.5

MARK = re.compile(r"\[\^(\d+)\]")
DEF = re.compile(r"^\[\^(\d+)\]:\s*(.+)$", re.M)

W_NS = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"


# ---------- 字体：同时设置 ascii/hAnsi/eastAsia 三个槽位 ----------

def set_font(run, name, size=None, bold=False):
    run.font.name = name
    run.font.size = size
    run.font.bold = bold
    rPr = run._element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(slot), name)


def style_font(style, name, size, bold=False):
    style.font.name = name
    style.font.size = size
    style.font.bold = bold
    rPr = style.element.get_or_add_rPr()
    rFonts = rPr.find(qn("w:rFonts"))
    if rFonts is None:
        rFonts = OxmlElement("w:rFonts")
        rPr.insert(0, rFonts)
    for slot in ("w:ascii", "w:hAnsi", "w:eastAsia", "w:cs"):
        rFonts.set(qn(slot), name)


def add_footnote_ref(paragraph, fid):
    """在段落末尾插入脚注引用标记（正文中的上标数字）。"""
    run = paragraph.add_run()
    rPr = run._element.get_or_add_rPr()
    rStyle = OxmlElement("w:rStyle")
    rStyle.set(qn("w:val"), "FootnoteReference")
    rPr.append(rStyle)
    ref = OxmlElement("w:footnoteReference")
    ref.set(qn("w:id"), str(fid))
    run._element.append(ref)


def first_line_indent_chars(paragraph, chars=2):
    """用 Word 的"字符"单位设置首行缩进，避免字号变化后缩进失真。"""
    pPr = paragraph._p.get_or_add_pPr()
    ind = pPr.find(qn("w:ind"))
    if ind is None:
        ind = OxmlElement("w:ind")
        pPr.append(ind)
    ind.set(qn("w:firstLineChars"), str(chars * 100))
    ind.set(qn("w:firstLine"), str(240 * chars))


TOKEN = re.compile(r"(\*\*.+?\*\*|\[\^\d+\])")


def add_text_with_marks(paragraph, text):
    """写入正文：解析行内 **加粗** 与 [^n] 脚注标记。

    早期版本只处理脚注，结果正文里的 ** 被原样打了出来——这类缺陷只有渲染成
    图片才看得见，纯文本检查发现不了。
    """
    for tok in TOKEN.split(text):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**") and len(tok) > 4:
            set_font(paragraph.add_run(tok[2:-2]), BODY_FONT, BODY_SIZE, bold=True)
        elif MARK.fullmatch(tok):
            add_footnote_ref(paragraph, int(MARK.fullmatch(tok).group(1)))
        else:
            set_font(paragraph.add_run(tok), BODY_FONT, BODY_SIZE)


# ---------- 脚注部件 ----------

def footnote_paragraph_xml(fid, text):
    """生成一条脚注：脚注编号标记 + 正文。内联加粗标记 **x** 被还原为加粗 run。"""
    runs = []
    pos = 0
    for m in re.finditer(r"\*\*(.+?)\*\*", text):
        if m.start() > pos:
            runs.append((text[pos:m.start()], False))
        runs.append((m.group(1), True))
        pos = m.end()
    if pos < len(text):
        runs.append((text[pos:], False))
    if not runs:
        runs = [(text, False)]

    body = []
    for seg, bold in runs:
        seg = seg.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        rpr = f'<w:rPr><w:rFonts w:ascii="{BODY_FONT}" w:hAnsi="{BODY_FONT}" w:eastAsia="{BODY_FONT}" w:cs="{BODY_FONT}"/><w:sz w:val="{FOOT_SIZE.pt * 2:.0f}"/><w:szCs w:val="{FOOT_SIZE.pt * 2:.0f}"/>'
        if bold:
            rpr += "<w:b/>"
        rpr += "</w:rPr>"
        body.append(f'<w:r>{rpr}<w:t xml:space="preserve">{seg}</w:t></w:r>')

    return (
        f'<w:footnote w:id="{fid}">'
        f'<w:p><w:pPr><w:pStyle w:val="FootnoteText"/></w:pPr>'
        f'<w:r><w:rPr><w:rStyle w:val="FootnoteReference"/></w:rPr><w:footnoteRef/></w:r>'
        f'<w:r><w:rPr><w:rFonts w:ascii="{BODY_FONT}" w:hAnsi="{BODY_FONT}" w:eastAsia="{BODY_FONT}"/>'
        f'<w:sz w:val="{FOOT_SIZE.pt * 2:.0f}"/></w:rPr><w:t xml:space="preserve"> </w:t></w:r>'
        + "".join(body)
        + "</w:p></w:footnote>"
    )


def build_footnotes_xml(defs):
    sep = (
        '<w:footnote w:type="separator" w:id="-1"><w:p><w:pPr>'
        '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:r><w:separator/></w:r></w:p></w:footnote>'
        '<w:footnote w:type="continuationSeparator" w:id="0"><w:p><w:pPr>'
        '<w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
        '<w:r><w:continuationSeparator/></w:r></w:p></w:footnote>'
    )
    notes = "".join(footnote_paragraph_xml(i, defs[i]) for i in sorted(defs))
    return (
        '<?xml version="1.0" encoding="UTF-8" standalone="yes"?>\n'
        f'<w:footnotes xmlns:w="{W_NS}">' + sep + notes + "</w:footnotes>"
    )


FOOTNOTE_STYLES = (
    '<w:style w:type="character" w:styleId="FootnoteReference">'
    '<w:name w:val="footnote reference"/><w:rPr><w:vertAlign w:val="superscript"/></w:rPr></w:style>'
    '<w:style w:type="paragraph" w:styleId="FootnoteText">'
    '<w:name w:val="footnote text"/>'
    '<w:pPr><w:spacing w:after="0" w:line="240" w:lineRule="auto"/></w:pPr>'
    f'<w:rPr><w:rFonts w:ascii="{BODY_FONT}" w:hAnsi="{BODY_FONT}" w:eastAsia="{BODY_FONT}"/>'
    f'<w:sz w:val="{FOOT_SIZE.pt * 2:.0f}"/><w:szCs w:val="{FOOT_SIZE.pt * 2:.0f}"/></w:rPr></w:style>'
)


def inject_parts(docx_path, defs):
    """把脚注部件、内容类型、关系项、样式注入已保存的 .docx。"""
    tmp = docx_path.with_suffix(".tmp.docx")
    with zipfile.ZipFile(docx_path) as zin:
        names = zin.namelist()
        data = {n: zin.read(n) for n in names}

    def dec(b):
        return b.decode("utf-8")

    # 1) 内容类型
    ct = dec(data["[Content_Types].xml"])
    if "footnotes+xml" not in ct:
        ct = ct.replace(
            "</Types>",
            '<Override PartName="/word/footnotes.xml" ContentType="application/vnd.openxmlformats-'
            'officedocument.wordprocessingml.footnotes+xml"/></Types>',
        )
    data["[Content_Types].xml"] = ct.encode("utf-8")

    # 2) 关系项
    rels = dec(data["word/_rels/document.xml.rels"])
    if "footnotes.xml" not in rels:
        used = re.findall(r'Id="rId(\d+)"', rels)
        rid = f"rId{max((int(x) for x in used), default=0) + 1}"
        rels = rels.replace(
            "</Relationships>",
            f'<Relationship Id="{rid}" Type="http://schemas.openxmlformats.org/officeDocument/2006/'
            f'relationships/footnotes" Target="footnotes.xml"/></Relationships>',
        )
    data["word/_rels/document.xml.rels"] = rels.encode("utf-8")

    # 3) 样式
    styles = dec(data["word/styles.xml"])
    if "FootnoteReference" not in styles:
        styles = styles.replace("</w:styles>", FOOTNOTE_STYLES + "</w:styles>")
    data["word/styles.xml"] = styles.encode("utf-8")

    # 4) 脚注部件本体
    data["word/footnotes.xml"] = build_footnotes_xml(defs).encode("utf-8")

    with zipfile.ZipFile(tmp, "w", zipfile.ZIP_DEFLATED) as zout:
        for n in names:
            zout.writestr(n, data[n])
        zout.writestr("word/footnotes.xml", data["word/footnotes.xml"])
    shutil.move(str(tmp), str(docx_path))


# ---------- 主流程 ----------

def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    a = ap.parse_args()

    text = pathlib.Path(a.src).read_text(encoding="utf-8")
    if "\n## 脚注" in text:
        body, defs_raw = text.split("\n## 脚注", 1)
    else:
        body, defs_raw = text, ""
    defs = {int(k): v.strip() for k, v in DEF.findall(defs_raw)}

    doc = Document()

    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
    sec.top_margin = sec.bottom_margin = Cm(2.54)
    sec.left_margin = sec.right_margin = Cm(3.18)

    normal = doc.styles["Normal"]
    style_font(normal, BODY_FONT, BODY_SIZE)
    normal.paragraph_format.line_spacing_rule = WD_LINE_SPACING.MULTIPLE
    normal.paragraph_format.line_spacing = LINE_SPACING
    normal.paragraph_format.space_after = Pt(0)

    for name in ("Heading 1", "Heading 2", "Heading 3"):
        try:
            st = doc.styles[name]
            style_font(st, HEAD_FONT, Pt(14 if name == "Heading 1" else 12), bold=True)
            # python-docx 默认 Heading 样式带主题蓝，中文论文标题须为黑色
            st.font.color.rgb = RGBColor(0, 0, 0)
        except KeyError:
            pass

    used = set()
    for raw in body.splitlines():
        line = raw.strip()
        if not line or line == "---":
            continue

        if line.startswith("# "):                      # 文题
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            p.paragraph_format.space_after = Pt(18)
            set_font(p.add_run(line[2:]), HEAD_FONT, Pt(18), bold=True)

        elif line.startswith("### "):                  # 节标题
            p = doc.add_paragraph(style="Heading 2")
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(6)
            set_font(p.add_run(line[4:]), HEAD_FONT, Pt(12), bold=True)

        elif line.startswith("## "):                   # 章标题
            p = doc.add_paragraph(style="Heading 1")
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(8)
            set_font(p.add_run(line[3:]), HEAD_FONT, Pt(14), bold=True)

        else:                                          # 正文段落
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            if line.startswith("**摘要**") or line.startswith("**关键词**"):
                # 摘要与关键词不缩进，标签加粗
                label, _, rest = line.partition("**：")
                set_font(p.add_run(label.replace("**", "") + "："), HEAD_FONT, BODY_SIZE, bold=True)
                add_text_with_marks(p, rest)
            else:
                first_line_indent_chars(p, 2)
                add_text_with_marks(p, line)

    for n in MARK.findall(body):
        used.add(int(n))

    doc.save(a.dst)
    inject_parts(pathlib.Path(a.dst), defs)

    missing = sorted(set(defs) - used)
    print(f"正文段落已写入 {a.dst}")
    print(f"脚注定义 {len(defs)} 条，正文实际引用 {len(used)} 条")
    if missing:
        print(f"警告：以下脚注有定义但正文未引用：{missing}")


if __name__ == "__main__":
    main()
